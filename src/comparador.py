from collections import defaultdict
from pathlib import Path
import os
import sys
from xml.sax.saxutils import escape

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


COLUMNAS_REQUERIDAS = {
    "Producto",
    "Proveedor",
    "Precio",
    "PlazoEntrega",
    "CondicionesPago",
}


def obtener_raiz_aplicacion():
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parents[2]
    return Path(__file__).resolve().parents[1]


def obtener_archivo_inicial():
    if getattr(sys, "frozen", False):
        carpeta_temporal = getattr(sys, "_MEIPASS", None)
        if carpeta_temporal:
            return Path(carpeta_temporal) / "data" / "raw" / "cotizaciones.xlsx"
    return obtener_raiz_aplicacion() / "data" / "raw" / "cotizaciones.xlsx"


def calcular_puntajes(df):
    columnas_faltantes = COLUMNAS_REQUERIDAS.difference(df.columns)
    if columnas_faltantes:
        faltantes = ", ".join(sorted(columnas_faltantes))
        raise ValueError(f"El archivo no contiene las columnas requeridas: {faltantes}.")
    if df.empty:
        raise ValueError("El archivo de cotizaciones no contiene registros.")

    max_precio = df["Precio"].max()
    max_tiempo = df["PlazoEntrega"].max()
    max_pago = df["CondicionesPago"].max()
    if max_precio == 0 or max_tiempo == 0 or max_pago == 0:
        raise ValueError("Los valores máximos de precio, plazo y pago deben ser distintos de cero.")

    df = df.copy()
    df["Score_Precio"] = (1 - (df["Precio"] / max_precio)) * 60
    df["Score_Tiempo"] = (1 - (df["PlazoEntrega"] / max_tiempo)) * 25
    df["Score_Pago"] = (df["CondicionesPago"] / max_pago) * 15
    df["Score_Total"] = (
        df["Score_Precio"] + df["Score_Tiempo"] + df["Score_Pago"]
    )
    return df


def agregar_grafico_puntajes(ruta_archivo):
    libro = load_workbook(ruta_archivo)
    hoja = libro.active
    encabezados = {
        str(celda.value).strip(): celda.column
        for celda in hoja[1]
        if celda.value is not None
    }
    if "Proveedor" not in encabezados or "Score_Total" not in encabezados:
        raise ValueError("No se encontraron las columnas 'Proveedor' y 'Score_Total'.")

    puntajes = defaultdict(lambda: [0.0, 0])
    for fila in range(2, hoja.max_row + 1):
        proveedor = hoja.cell(fila, encabezados["Proveedor"]).value
        puntaje = hoja.cell(fila, encabezados["Score_Total"]).value
        if proveedor is None or puntaje is None:
            continue

        nombre = str(proveedor).strip()
        if not nombre:
            continue
        if not isinstance(puntaje, (int, float)):
            raise ValueError(
                f"Puntaje no numérico para el proveedor '{nombre}' en la fila {fila}."
            )
        puntajes[nombre][0] += puntaje
        puntajes[nombre][1] += 1

    if not puntajes:
        raise ValueError("No hay puntajes por proveedor para graficar.")

    columna_proveedor = 11
    columna_puntaje = 12
    hoja.cell(1, columna_proveedor, "Proveedor")
    hoja.cell(1, columna_puntaje, "Puntaje promedio")
    for fila, (proveedor, (total, cantidad)) in enumerate(puntajes.items(), start=2):
        hoja.cell(fila, columna_proveedor, proveedor)
        hoja.cell(fila, columna_puntaje, total / cantidad)

    grafico = BarChart()
    grafico.type = "bar"
    grafico.style = 10
    grafico.title = "Puntaje promedio por proveedor"
    grafico.x_axis.title = "Puntaje promedio"
    grafico.y_axis.title = "Proveedor"
    grafico.y_axis.scaling.orientation = "maxMin"
    datos = Reference(
        hoja,
        min_col=columna_puntaje,
        min_row=1,
        max_row=len(puntajes) + 1,
    )
    categorias = Reference(
        hoja,
        min_col=columna_proveedor,
        min_row=2,
        max_row=len(puntajes) + 1,
    )
    grafico.add_data(datos, titles_from_data=True)
    grafico.set_categories(categorias)
    grafico.height = 8
    grafico.width = 16
    hoja.add_chart(grafico, "N2")
    libro.save(ruta_archivo)


def generar_pdf_top_proveedores(df_resultados, ruta_pdf):
    top_5 = df_resultados.sort_values(
        by="Score_Total", ascending=False
    ).head(5)
    documento = SimpleDocTemplate(str(ruta_pdf), pagesize=letter)
    contenido = []
    estilos = getSampleStyleSheet()
    contenido.append(
        Paragraph(
            "<b>Reporte Oficial: Top 5 Proveedores</b>",
            estilos["Title"],
        )
    )
    contenido.append(Spacer(1, 15))

    for _, fila in top_5.iterrows():
        texto = (
            f"<b>Producto:</b> {escape(str(fila['Producto']))}"
            f"<br/><b>Proveedor recomendado:</b> {escape(str(fila['Proveedor']))}"
            f"<br/><b>Puntaje total:</b> {fila['Score_Total']:.2f}"
        )
        contenido.append(Paragraph(texto, estilos["Normal"]))
        contenido.append(Spacer(1, 10))

    documento.build(contenido)


def procesar_cotizaciones(ruta_entrada):
    ruta_salida = obtener_raiz_aplicacion() / "output"
    ruta_salida.mkdir(parents=True, exist_ok=True)
    df = calcular_puntajes(pd.read_excel(ruta_entrada))

    ruta_excel = ruta_salida / "resultado_cotizaciones.xlsx"
    ruta_pdf = ruta_salida / "top_5_proveedores.pdf"
    df.to_excel(ruta_excel, index=False)
    agregar_grafico_puntajes(ruta_excel)
    generar_pdf_top_proveedores(df, ruta_pdf)
    return ruta_excel, ruta_pdf, df


if __name__ == "__main__":
    archivo_entrada = obtener_archivo_inicial()
    print(f"Procesando archivo: {archivo_entrada}")
    ruta_excel, ruta_pdf, _ = procesar_cotizaciones(archivo_entrada)
    print("¡Proceso completado con éxito!")
    print(f"Excel generado en: {ruta_excel}")
    print(f"PDF generado en: {ruta_pdf}")